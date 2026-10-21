#!/usr/bin/env python3
# 作者与使用声明 / Authorship and use: see NOTICE.md.
# 代码由本仓库作者编写；实际应用风险由使用者自行评估并承担（以适用法律为限）。
# Written by the repository author; practical use is at the user's risk, subject to applicable law.
# 原始 PDF 和数据来自互联网，权利归原权利人 / Original PDF and data are third-party materials.
"""回放真实录波，联动继电器与理想断路器，导出 XLSX、PNG 和无文字阻抗 GIF。

python simulate_with_relay.py
python simulate_with_relay.py --breaker-delay 0.08
python simulate_with_relay.py --gif-only

注意：这是录波驱动的简化闭环演示，不是完整电力网络电磁暂态仿真。
断开后，线路 CT 电流为零，电源侧 PT 电压取故障前基波外推值。
"""

from __future__ import annotations

import argparse
import cmath
import json
import math
import statistics
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

import distance_relay as relay
from data import resolve_font_path


def simulate(cfg_path: Path, delay: float, characteristic: str = 'mho') -> dict:
    if not math.isfinite(delay) or delay < 0:
        raise ValueError('断路器延迟必须是非负有限数。')
    if characteristic != 'mho':
        raise ValueError('此版本只采用正方向 mho 圆。')
    cfg = relay.parse_cfg(cfg_path)
    times, original = relay.read_dat(cfg_path.with_suffix('.dat'), cfg)
    if len(cfg.channels) != 6 or any(c.unit.strip().upper() != 'A' for c in cfg.channels[:3]):
        raise ValueError('需要本作业的 IA、IB、IC、UA、UB、UC 六通道录波。')
    event = 'Event1' if 'Event1' in cfg_path.stem else 'Event2'
    loop = '自动'
    settings = relay.build_settings(characteristic)
    n = round(cfg.sample_rate_hz / cfg.frequency_hz)
    # 采用故障前最初一个周波估计各相电源侧电压；相位与时间原点一致。
    source_voltage = [
        relay.one_cycle_phasor([r[c] for r in original[:n]], times[:n], cfg.frequency_hz)
        for c in range(3, 6)
    ]
    bank = relay.RelayBank()
    state = bank.state
    all_loops = {name: [] for name in relay.LOOPS}
    modeled: list[list[float]] = []
    rows = []
    open_time = None
    scheduled_open_time = None
    for k, (t, raw) in enumerate(zip(times, original)):
        if state.trip_time_s is not None and scheduled_open_time is None:
            scheduled_open_time = state.trip_time_s + delay
        opened = scheduled_open_time is not None and t >= scheduled_open_time
        if opened:
            if open_time is None:
                open_time = t
            rotation = cmath.exp(1j * 2 * math.pi * cfg.frequency_hz * t)
            values = [0.0, 0.0, 0.0] + [
                math.sqrt(2) * (u * rotation).real for u in source_voltage
            ]
        else:
            values = list(raw)
        modeled.append(values)
        result = None
        if k >= n - 1:
            first = k - n + 1
            phasors = [
                relay.one_cycle_phasor([r[c] for r in modeled[first:k+1]], times[first:k+1], cfg.frequency_hz)
                for c in range(6)
            ]
            loop, result, loop_results = bank.step(settings, t, times[k]-times[k-1], phasors)
            for name, item in loop_results.items():
                z = item.impedance
                all_loops[name].append([t,None if z is None else z.real,None if z is None else z.imag,
                                       item.in_zone1,item.in_zone2,item.in_zone3,item.zone2_timer_s,item.zone3_timer_s,
                                       item.direction,item.z2_projection])
            # 零延时断路器也只在收到命令后才改变样本，不提前修改测量。
        else:
            for history in all_loops.values():history.append([t,None,None,False,False,False,0.,0.,'窗口未满',None])
        if opened:
            stage = '断路器已断开'
        elif state.trip_latched:
            stage = '已发跳闸命令，等待断开'
        else:
            stage = '断路器闭合'
        z = result.impedance if result else None
        rows.append([
            k+1, round(t*1e6), t, *values,
            int(state.trip_latched), int(not opened), stage,
            result.region if result else '一周波窗口未满',
            result.zone2_timer_s if result else 0.0,
            result.zone3_timer_s if result else 0.0,
            z.real if z is not None else None,
            z.imag if z is not None else None,
            loop, result.direction if result else '窗口未满',
            result.z2_projection if result else None,
        ])
    # 回归检查：断开前数据保持原样，断开后电流为零，不出现提前切除。
    for t, old, new in zip(times, original, modeled):
        if open_time is None or t < open_time:
            assert old == new
        else:
            assert new[:3] == [0.0, 0.0, 0.0]
    if open_time is not None:
        assert open_time >= state.trip_time_s + delay - 1e-12
        assert open_time - (state.trip_time_s + delay) < 1.1 / cfg.sample_rate_hz
    directions = [r[18] for r in rows if r[18] in ('正向', '反向')]
    q_values = [r[19] for r in rows if r[19] is not None]
    dominant_direction = max(set(directions), key=directions.count) if directions else '未判定'
    fault_loops = [r[17] for r in rows if r[18] in ('正向', '反向')]
    loop = bank.trip_loop or (max(set(fault_loops), key=fault_loops.count) if fault_loops else loop)
    return {
        'event': event, 'loop': loop, 'frequency': cfg.frequency_hz,
        'sample_rate': cfg.sample_rate_hz,
        'characteristic': characteristic, 'breaker_delay': delay,
        'trip_time': state.trip_time_s, 'trip_zone': state.trip_zone,
        'open_time': open_time, 'rows': rows,
        'direction': dominant_direction,
        'direction_projection_median': statistics.median(q_values) if q_values else None,
        'direction_deadband': relay.DIRECTION_DEADBAND_OHM,
        'all_loops': all_loops,
        'trip_loop': bank.trip_loop,
        'raw_rows': [[k+1, round(t*1e6), t, *r] for k, (t, r) in enumerate(zip(times, original))],
        'headers': ['采样序号','时间戳(μs)','时间(s)','I_A(A)','I_B(A)','I_C(A)',
                    'U_AN(V)','U_BN(V)','U_CN(V)','跳闸命令(1为发出)','断路器闭合(1为闭合)',
                    '设备状态','当前动作区','II段计时(s)','III段计时(s)','Z实部(Ω)','Z虚部(Ω)',
                    '显示测量回路','32Q方向','负序方向投影(Ω)'],
        'zone_reaches': [abs(settings.zone1), abs(settings.zone2), abs(settings.zone3)],
        'zone_vectors': [[z.real,z.imag] for z in [settings.zone1,settings.zone2,settings.zone3]],
        'zone_centers': [[z.real/2,z.imag/2] for z in [settings.zone1,settings.zone2,settings.zone3]],
        'zone_radii': [abs(z)/2 for z in [settings.zone1,settings.zone2,settings.zone3]],
    }


def plot_event(payload: dict, output: Path) -> None:
    """保留 data 文件夹中上下两幅三相波形的形式，添加两个事件标记。"""
    image = Image.new('RGB', (1600, 1060), 'white')
    draw = ImageDraw.Draw(image)
    font_path = resolve_font_path()
    font = ImageFont.truetype(str(font_path), 19)
    small = ImageFont.truetype(str(font_path), 15)
    title = ImageFont.truetype(str(font_path), 25)
    event = payload['event']
    draw.text((70, 18), f'{event} — 继电器投入后三相波形（简化断路器仿真）', fill='#172B4D', font=title)
    trip = payload['trip_time']
    opened = payload['open_time']
    trip_text = '未发出跳闸命令' if trip is None else f"{payload['trip_zone']}跳闸命令：{trip:.6f} s"
    open_text = '未断开' if opened is None else f'断路器断开：{opened:.6f} s'
    draw.text((70, 57), f"{trip_text}    {open_text}    断路器延迟={payload['breaker_delay']*1000:.0f} ms", fill='#333333', font=font)
    draw.text((70, 89), f"动作/诊断回路 {payload['loop']}；六回路独立并行；32Q {payload['direction']}。模型：开断后电流为零，电源侧电压外推。", fill='#555555', font=small)
    rows = payload['rows']
    start_time, end_time = rows[0][2], rows[-1][2]
    left, right = 140, 1540
    colors = ['#d62728', '#2ca02c', '#1f77b4']
    for top, bottom, columns, label, names in [
        (190,490,[3,4,5],'电流 I (A)',['I_A','I_B','I_C']),
        (650,950,[6,7,8],'电压 U (V)',['U_AN','U_BN','U_CN']),
    ]:
        maximum = max(abs(row[c]) for row in rows for c in columns) * 1.08 or 1
        def xy(t, v):
            return (left+(t-start_time)/(end_time-start_time)*(right-left), (top+bottom)/2-v/maximum*(bottom-top)/2)
        draw.rectangle((left,top,right,bottom), outline='#444444')
        draw.text((left,top-35), f'{event} — 三相' + ('电流波形' if columns[0]==3 else '电源侧对地电压波形'), fill='#222222',font=font)
        # 中文竖向轴标题
        label_img = Image.new('RGBA',(170,34),(255,255,255,0))
        ImageDraw.Draw(label_img).text((2,2),label,fill='#222222',font=font)
        label_img = label_img.rotate(90,expand=True)
        image.paste(label_img,(12,(top+bottom-label_img.height)//2),label_img)
        for j in range(7):
            t=start_time+(end_time-start_time)*j/6
            x=xy(t,0)[0]
            draw.line((x,top,x,bottom),fill='#dddddd')
            draw.text((x-19,bottom+8),f'{t:.2f}',font=small,fill='#333333')
        for j in range(5):
            val=-maximum+2*maximum*j/4
            y=xy(start_time,val)[1]
            draw.line((left,y,right,y),fill='#dddddd')
            text=f'{val:.0f}'
            draw.text((left-12-draw.textlength(text,font=small),y-9),text,fill='#333333',font=small)
        for c,color,name in zip(columns,colors,names):
            draw.line([xy(row[2],row[c]) for row in rows],fill=color,width=1)
            pos=right-290+(c-columns[0])*100
            draw.line((pos,top-21,pos+20,top-21),fill=color,width=3)
            draw.text((pos+25,top-32),name,fill='#333333',font=small)
        for t,color in [(trip,'#8e44ad'),(opened,'#d97706')]:
            if t is not None:
                x=xy(t,0)[0]
                for y in range(top,bottom,12):
                    draw.line((x,y,x,min(y+6,bottom)),fill=color,width=2)
        draw.text((750,bottom+38),'时间 (s)',font=font,fill='#333333')
    draw.text((140,1020),'紫色虚线：跳闸命令    橙色虚线：断路器断开    电压测点：电源侧；电流测点：线路支路',fill='#555555',font=small)
    image.save(output)


def export_excel(payload: dict, output: Path) -> None:
    """Public Python dependency only; export values, not spreadsheet formulas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    main = wb.active
    main.title = '继电器投入后波形'
    raw = wb.create_sheet('原始录波')
    main.merge_cells('A1:Q1')
    main['A1'] = f"{payload['event']} 继电器与理想断路器仿真"
    main['A1'].font = Font(name='宋体', size=16, bold=True)
    main.row_dimensions[1].height = 30
    metadata = [
        ['项目', '本次计算结果'],
        ['测量回路', payload['loop']],
        ['动作特性', '正方向 mho 圆'],
        ['跳闸段别', payload['trip_zone']],
        ['跳闸命令时间(s)', payload['trip_time']],
        ['断路器延迟(s)', payload['breaker_delay']],
        ['实际断开时间(s)', payload['open_time']],
    ]
    for row_index, values in enumerate(metadata, 4):
        for col, value in enumerate(values, 1):
            main.cell(row_index, col, value)
    notes = [
        '仿真假设', '断开前沿用原录波，断开后线路三相电流为零。',
        '电压测点为电源侧，断开后按故障前基波外推。',
        '不模拟电弧、重合闸、过电压和完整电网暂态。',
        '未动作时不切除原始故障；原始录波仍可能自行恢复。',
        f"采样率 {payload['sample_rate']:.3f} Hz，基波 {payload['frequency']:.1f} Hz",
        'I/II/III 半径(Ω)：' + '/'.join(f'{v:.3f}' for v in payload['zone_radii']),
    ]
    for row_index, note in enumerate(notes, 4):
        main.cell(row_index, 4, note)
    for row_index, values in enumerate([payload['headers'], *payload['rows']], 12):
        for col, value in enumerate(values, 1):
            main.cell(row_index, col, value)
    raw['A1'] = f"原始录波来源：EE271_DistanceProtection_{payload['event']}.cfg / .dat"
    raw.merge_cells('A1:I1')
    for row_index, values in enumerate([payload['headers'][:9], *payload['raw_rows']], 3):
        for col, value in enumerate(values, 1):
            raw.cell(row_index, col, value)
    loop_sheets = []
    for name, history in payload['all_loops'].items():
        ws = wb.create_sheet(name+'回路')
        ws.append(['时间(s)','R(Ω)','X(Ω)','I段启动','II段启动','III段启动',
                   'II段计时(s)','III段计时(s)','32Q方向','负序方向投影(Ω)'])
        for record in history:ws.append(record)
        loop_sheets.append((ws,1,10))
    head_fill = PatternFill('solid', fgColor='243E60')
    stripe_fill = PatternFill('solid', fgColor='F2F5F8')
    thin = Side(style='thin', color='D9D9D9')
    for ws, header_row, columns in [(main, 12, len(payload['headers'])), (raw, 3, 9), *loop_sheets]:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f'A{header_row + 1}'
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(columns)}{ws.max_row}'
        for col in range(1, columns + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name='宋体', size=10)
                cell.alignment = Alignment(vertical='center')
                if cell.row >= header_row:
                    if cell.row == header_row:
                        cell.fill = head_fill
                        cell.font = Font(name='宋体', size=10, bold=True, color='FFFFFF')
                        cell.alignment = Alignment(vertical='center', wrap_text=True)
                        cell.border = Border(bottom=thin)
                    else:
                        if cell.row % 2 == 0:
                            cell.fill = stripe_fill
                        if isinstance(cell.value, float):
                            cell.number_format = '0.000000'
            ws.row_dimensions[row[0].row].height = 20
        ws.row_dimensions[header_row].height = 32
    main['A1'].font = Font(name='宋体', size=16, bold=True)
    main.column_dimensions['A'].width = 26
    main.column_dimensions['B'].width = 29
    main.column_dimensions['L'].width = 32
    for row_index in range(8, 11):
        main.cell(row_index, 2).number_format = '0.000000'
    wb.save(output)


def plot_impedance_gif(payload: dict, output: Path, start=.98, end=1.05, frame_count=120):
    """Compatibility wrapper: six independent loops, never the old two-panel view."""
    return plot_all_loops(payload, output, start, end, frame_count, True)


def plot_all_loops(payload: dict, output: Path, start=.98, end=1.05, frame_count=120, animated=False):
    """参考图式六宫格：同一定值、同一坐标比例，六条独立回路的实际轨迹。"""
    width,height = (1200,800) if animated else (1800,1200)
    base=Image.new('RGB',(width,height),'white');d=ImageDraw.Draw(base)
    font=None if animated else ImageFont.truetype(str(resolve_font_path()),26)
    small=None if animated else ImageFont.truetype(str(resolve_font_path()),20)
    if not animated:
        d.text((50,12),payload['event']+' 六测量回路三段方向阻抗圆及实际轨迹',fill='black',font=font)
        d.text((50,height-42),'蓝色 I 段   红色 II 段   绿色 III 段   黑线：真实测量轨迹   时间窗 0.98–1.05 s',fill='black',font=small)
    xmin,xmax,ymin,ymax=-130,280,-130,210
    panels=[];colors=['#187fb5','#d4452d','#159765']
    for pos,name in enumerate(relay.LOOPS):
        col,row=pos%3,pos//3
        bx=col*width/3;by=row*(height-80)/2+50
        availw=width/3-105;availh=(height-100)/2-90
        scale=min(availw/(xmax-xmin),availh/(ymax-ymin))
        w,h=round((xmax-xmin)*scale),round((ymax-ymin)*scale)
        left=round(bx+70+(availw-w)/2);top=round(by+30)
        layer=Image.new('RGB',(w+1,h+1),'white');g=ImageDraw.Draw(layer)
        def xy(z,scale=scale):return ((z.real-xmin)*scale,(ymax-z.imag)*scale)
        for v in [-100,0,100,200]:
            x,y=xy(complex(v,v));g.line((x,0,x,h),fill='#e0e0e0');g.line((0,y,w,y),fill='#e0e0e0')
            if not animated:
                d.text((left+x-18,top+h+5),str(v),font=small,fill='#333333')
                d.text((left-48,top+y-12),str(v),font=small,fill='#333333')
        g.rectangle((0,0,w,h),outline='#999999')
        for center,radius,color in reversed(list(zip(payload['zone_centers'],payload['zone_radii'],colors))):
            z=complex(*center);g.ellipse((*xy(z+complex(-radius,radius)),*xy(z+complex(radius,-radius))),outline=color,width=2 if animated else 3)
        if not animated:
            d.text((left+w/2-20,top-35),name,font=font,fill='black')
            d.text((left+w/2-25,top+h+32),'R / Ω',font=small,fill='black')
            d.text((left+3,top-25),'X / Ω',font=small,fill='black')
        points=[(r[0],complex(r[1],r[2])) for r in payload['all_loops'][name] if r[1] is not None and start<=r[0]<=end
                and (payload['open_time'] is None or r[0]<payload['open_time'])]
        panels.append((layer,(left,top),xy,points))
    def frame(t):
        im=base.copy()
        for layer,origin,xy,points in panels:
            imlayer=layer.copy();g=ImageDraw.Draw(imlayer);trail=[xy(z) for ti,z in points if ti<=t]
            if len(trail)>1:g.line(trail,fill='black',width=2)
            if trail:
                x,y=trail[-1];g.ellipse((x-3,y-3,x+3,y+3),fill='black')
            im.paste(imlayer,origin)
        return im
    if not animated:frame(end).save(output);return
    frames=[frame(start+(end-start)*k/(frame_count-1)) for k in range(frame_count)]
    frames[0].save(output,save_all=True,append_images=frames[1:],duration=[600]+[50]*(frame_count-2)+[1600],loop=0,disposal=2,optimize=True)


def main() -> None:
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir',type=Path,default=here/'inputs'/'comtrade')
    parser.add_argument('--output-dir',type=Path,default=here/'outputs'/'simulation')
    parser.add_argument('--breaker-delay',type=float,default=0.05,help='断路器动作延迟(s)，默认0.05')
    parser.add_argument('--event',choices=['1','2','all'],default='all')
    parser.add_argument('--gif-only',action='store_true',help='只生成无文字阻抗轨迹 GIF')
    parser.add_argument('--gif-start',type=float,default=.98)
    parser.add_argument('--gif-end',type=float,default=1.05)
    parser.add_argument('--gif-frames',type=int,default=120)
    args=parser.parse_args()
    if not math.isfinite(args.breaker_delay) or args.breaker_delay<0:
        parser.error('断路器延迟必须是非负有限数。')
    if not all(math.isfinite(v) for v in [args.gif_start,args.gif_end]) or args.gif_start>=args.gif_end:
        parser.error('GIF 起止时刻必须是递增的有限数。')
    if not 2 <= args.gif_frames <= 600:
        parser.error('GIF 帧数须在 2～600 之间。')
    args.output_dir.mkdir(parents=True,exist_ok=True)
    summaries=[]
    events=['Event1','Event2'] if args.event=='all' else [f'Event{args.event}']
    for event in events:
        item=simulate(args.input_dir/f'EE271_DistanceProtection_{event}.cfg',args.breaker_delay)
        if not args.gif_only:
            plot_event(item,args.output_dir/f'{event}_with_relay_waveforms.png')
            export_excel(item,args.output_dir/f'{event}_with_relay_data.xlsx')
            plot_all_loops(item,args.output_dir/f'{event}_RX_all_loops.png',args.gif_start,args.gif_end)
        plot_all_loops(item,args.output_dir/f'{event}_RX_trajectory.gif',args.gif_start,args.gif_end,args.gif_frames,True)
        summaries.append({key:item[key] for key in ['event','loop','characteristic','trip_time','trip_zone','open_time','breaker_delay','zone_reaches','zone_vectors','zone_centers','zone_radii','direction','direction_projection_median','direction_deadband','trip_loop']})
        print(f"{event}: 跳闸={item['trip_time']} s，断开={item['open_time']} s",flush=True)
    (args.output_dir/'summary.json').write_text(json.dumps(summaries,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'已生成仿真结果：{args.output_dir.resolve()}')


if __name__=='__main__':
    main()
